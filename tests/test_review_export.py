"""阶段七测试：人工审核与 Excel 导出。"""

import io
import uuid
from decimal import Decimal

import openpyxl

from tradepilot.app.models.quotation_item import QuotationItem, ReviewStatus


def make_task(client, name="审核导出测试") -> str:
    return client.post("/api/v1/tasks", json={"task_name": name}).json()["id"]


def seed_items(session_factory, task_id, n=3, confirmed=0):
    db = session_factory()
    try:
        for i in range(n):
            db.add(
                QuotationItem(
                    task_id=uuid.UUID(task_id),
                    supplier_name=f"供应商{i}",
                    product_name="不锈钢螺栓",
                    model="M8x30",
                    unit="件",
                    unit_price=Decimal("0.8") + Decimal(i) / 100,
                    currency="CNY",
                    quantity=Decimal("1000"),
                    review_status=ReviewStatus.CONFIRMED if i < confirmed else ReviewStatus.PENDING_REVIEW,
                )
            )
        db.commit()
    finally:
        db.close()


def run_match(client, task_id):
    return client.post(f"/api/v1/tasks/{task_id}/match").json()


# ---------- 审核概览 ----------


def test_review_summary_counts_and_issues(client, session_factory):
    task_id = make_task(client)
    db = session_factory()
    try:
        db.add(QuotationItem(task_id=uuid.UUID(task_id), product_name="缺价品"))  # 无单价无币种
        db.commit()
    finally:
        db.close()
    summary = client.get(f"/api/v1/tasks/{task_id}/review").json()
    assert summary["total_items"] == 1
    assert summary["pending_review"] == 1
    assert {i["issue"] for i in summary["issues"]} == {"缺少单价，无法参与比价", "缺少币种，需人工确认"}


def test_review_confirm_all_and_by_ids(client, session_factory):
    task_id = make_task(client)
    seed_items(session_factory, task_id, n=3)

    summary = client.get(f"/api/v1/tasks/{task_id}/review").json()
    ids = client.get(f"/api/v1/tasks/{task_id}/items").json()["items"]
    target_id = ids[0]["id"]

    # 按 ID 确认 1 条
    resp = client.post(
        f"/api/v1/tasks/{task_id}/review/confirm", json={"item_ids": [target_id]}
    ).json()
    assert resp["confirmed"] == 1

    # 全部确认
    resp = client.post(f"/api/v1/tasks/{task_id}/review/confirm", json={"all": True}).json()
    assert resp["confirmed"] == 2

    summary = client.get(f"/api/v1/tasks/{task_id}/review").json()
    assert summary["pending_review"] == 0
    assert summary["by_status"]["confirmed"] == 3


def test_review_confirm_unknown_task_404(client):
    resp = client.post(f"/api/v1/tasks/{uuid.uuid4()}/review/confirm", json={"all": True})
    assert resp.status_code == 404


# ---------- 导出 ----------


def test_export_xlsx_content(client, session_factory):
    task_id = make_task(client, "10月紧固件比价")
    seed_items(session_factory, task_id, n=2)
    run_match(client, task_id)

    resp = client.get(f"/api/v1/tasks/{task_id}/export")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["比价结果", "报价明细"]

    cmp_sheet = wb["比价结果"]
    assert cmp_sheet["A2"].value == "不锈钢螺栓"
    assert cmp_sheet["G2"].value == "供应商0"  # 最低价供应商（0.80 < 0.81）
    assert cmp_sheet["H2"].value == 0.8

    items_sheet = wb["报价明细"]
    assert items_sheet.max_row == 3  # 表头 + 2 条
    assert items_sheet["A2"].value == "供应商0"
    assert items_sheet["M2"].value == "待复核"

    # 导出后任务状态置 completed
    task = client.get(f"/api/v1/tasks/{task_id}").json()
    assert task["status"] == "completed"


def test_export_unknown_task_404(client):
    resp = client.get(f"/api/v1/tasks/{uuid.uuid4()}/export")
    assert resp.status_code == 404
